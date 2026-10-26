"""Verificacao de consistencia do estudo (criterios de aceite). Sai com codigo 1 se algo FALHAR.

Uso: python3 scripts/verificar.py   (requer LibreOffice para recalcular as formulas dos .xlsx)
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

import projeto as P

falhas, avisos, oks = [], [], []


def ok(cond, msg, aviso=False):
    (oks if cond else (avisos if aviso else falhas)).append(msg)


def sobrep(a, b):
    return min(a["x1"], b["x1"]) - max(a["x0"], b["x0"]) > 1e-6 and min(a["y1"], b["y1"]) - max(a["y0"], b["y0"]) > 1e-6


def dentro(a, x0, y0, x1, y1, tol=1e-6):
    return a["x0"] >= x0 - tol and a["y0"] >= y0 - tol and a["x1"] <= x1 + tol and a["y1"] <= y1 + tol


amb = P.ambientes()
res = P.resumo(amb)
hsh = P.hash_dados()

# 1. fechamento por pavimento (tolerancia 2%)
for p, r in res.items():
    d = abs(r["bruta"] / r["declarada"] - 1)
    ok(d <= 0.02, f"[areas] pav {p}: soma {r['bruta']:.2f} x declarada {r['declarada']:.2f} (desvio {d:.2%})")
    ok(abs(r["util"] + r["paredes"] - r["bruta"]) < 0.002, f"[areas] pav {p}: util + paredes = bruta")

# 2. sem sobreposicao; terreo cobre o lote inteiro
for p in "TS":
    sel = [a for a in amb if a["pav"] == p]
    for i, a in enumerate(sel):
        ok(dentro(a, 0, 0, P.LOTE_L, P.LOTE_P), f"[geom] {a['id']} dentro do lote")
        for b in sel[i + 1:]:
            ok(not sobrep(a, b), f"[geom] {a['id']} x {b['id']} sem sobreposicao")
tot_t = sum(a["area_bruta"] for a in amb if a["pav"] == "T")
ok(abs(tot_t - P.LOTE_L * P.LOTE_P) < 1e-6, f"[geom] ambientes do terreo somam a area do lote ({tot_t:.2f} m2)")

# 3. superior apoiado sobre o terreo (projecao dentro de ambientes solidos do terreo)
solT = [a for a in amb if a["pav"] == "T" and a["tipo"] in P.SOLIDOS]
for a in amb:
    if a["pav"] == "S" and a["tipo"] != "vazio":
        area_apoio = sum(max(0, min(a["x1"], b["x1"]) - max(a["x0"], b["x0"])) *
                         max(0, min(a["y1"], b["y1"]) - max(a["y0"], b["y0"])) for b in solT)
        ok(abs(area_apoio - a["area_bruta"]) < 1e-6, f"[geom] {a['id']} projeta sobre area construida do terreo")

# 4. recuo frontal (hipotese)
rf = P.PREMISSAS_LEGAIS["recuo_frontal_m"]
for a in amb:
    if a["tipo"] in P.SOLIDOS or a["id"] == "SACADA":
        ok(a["y0"] >= rf - 1e-6, f"[legal-hip] {a['id']} respeita recuo frontal de {rf} m")

# 5. janelas x divisas (Codigo Civil art. 1.301: 1,50 m)
for o in P.aberturas():
    if o["tipo"] not in ("janela", "porta_correr"):
        continue
    if o["eixo"] == "x":
        dist = min(o["coord"], P.LOTE_L - o["coord"])
    else:
        dist = P.LOTE_P - o["coord"] if o["para"] in ("EXT", "GOURMET", "JARDIM_F") else 99
    ok(dist >= P.PREMISSAS_LEGAIS["afast_janela_divisa_m"] - 1e-6,
       f"[legal] abertura {o['id']} a {dist:.2f} m da divisa que enxerga (>= 1,50 m)")

# 6. aberturas sobre parede dos ambientes indicados
idx = {(a["id"]): a for a in amb}
for o in P.aberturas():
    a = idx.get(o["de"])
    if not a:
        falhas.append(f"[aberturas] {o['id']}: ambiente {o['de']} inexistente"); continue
    if o["eixo"] == "y":
        on = abs(o["coord"] - a["y0"]) < 1e-6 or abs(o["coord"] - a["y1"]) < 1e-6
        rng = a["x0"] - 1e-6 <= o["a"] and o["b"] <= a["x1"] + 1e-6
    else:
        on = abs(o["coord"] - a["x0"]) < 1e-6 or abs(o["coord"] - a["x1"]) < 1e-6
        rng = a["y0"] - 1e-6 <= o["a"] and o["b"] <= a["y1"] + 1e-6
    ok(on and rng, f"[aberturas] {o['id']} esta na parede de {o['de']}")

# 7. garagem: 2 carros + 1 moto com circulacao
mob = {m["id"]: m for m in P.mobiliario()}
g = idx["GAR"]
gx0, gx1 = g["x0"] + g["ded"]["W"], g["x1"] - g["ded"]["E"]
gy0, gy1 = g["y0"] + g["ded"]["S"], g["y1"] - g["ded"]["N"]
veic = [mob["CAR_L"], mob["CAR_R"], mob["MOTO"]]
for v in veic:
    ok(dentro(v, gx0, gy0, gx1, gy1), f"[garagem] {v['id']} cabe na area util da garagem")
for i, a in enumerate(veic):
    for b in veic[i + 1:]:
        ok(not sobrep(a, b), f"[garagem] {a['id']} x {b['id']} sem sobreposicao")
folga_mot_esq = mob["CAR_L"]["x0"] - gx0
ok(folga_mot_esq >= 0.50, f"[garagem] folga porta do motorista carro 1 = {folga_mot_esq:.2f} m (>= 0,50)")
ok(folga_mot_esq >= 0.60, f"[garagem] folga carro 1 >= 0,60 m recomendada (atual {folga_mot_esq:.2f}) - abertura parcial", aviso=True)
corredor = mob["CAR_R"]["x0"] - mob["CAR_L"]["x1"]
ok(corredor >= 1.80, f"[garagem] faixa entre carros = {corredor:.2f} m (portas + moto)")
passagem = mob["MOTO"]["x0"] - mob["CAR_L"]["x1"]
ok(passagem >= 0.80, f"[garagem] passagem de pedestre ao lado da moto = {passagem:.2f} m (>= 0,80)")
porta = next(o for o in P.aberturas() if o["tipo"] == "porta_pivotante")
ok(porta["b"] > mob["CAR_L"]["x1"] + 0.6, "[garagem] porta principal acessivel pela faixa central")

# 8. lote: TO / TP (hipoteses)
lt = P.lote_resumo(amb)
ok(abs(lt["fechamento_lote"] - lt["area_lote"]) < 1e-6, "[lote] projecao + descobertos = area do lote")
ok(lt["to"] <= P.PREMISSAS_LEGAIS["taxa_ocupacao_max"],
   f"[legal-hip] TO {lt['to']:.1%} <= {P.PREMISSAS_LEGAIS['taxa_ocupacao_max']:.0%} (hipotese) - ver I-6 / pendencias",
   aviso=True)
ok(lt["tp"] >= P.PREMISSAS_LEGAIS["taxa_permeab_min"],
   f"[legal-hip] TP {lt['tp']:.1%} >= {P.PREMISSAS_LEGAIS['taxa_permeab_min']:.0%} (hipotese)")
ext_area = sum(e["area"] for e in P.externas())
ok(abs(ext_area - lt["descoberto_terreo"]) < 1e-6, "[lote] superficies externas = areas descobertas do terreo")
tot = sum(r["bruta"] for r in res.values())
ok(125 <= tot <= 140, f"[programa] area construida {tot:.2f} m2 dentro da meta 125-140 m2 - CONFLITO documentado em I-6",
   aviso=True)

# 9. mesma fonte de dados: plantas, 3D e tabela
for f in ("02_planta_terreo.svg", "03_planta_superior.svg"):
    ok(hsh in open(os.path.join(P.BASE, f), encoding="utf-8").read(), f"[fonte] {f} gerado com hash {hsh}")
dj = json.load(open(os.path.join(P.BASE, "14_modelo_3d", "dados_modelo.json"), encoding="utf-8"))
ok(dj["hash"] == hsh, "[fonte] 3D (dados_modelo.json) usa o mesmo hash")
html = open(os.path.join(P.BASE, "14_modelo_3d", "index.html"), encoding="utf-8").read()
ok(f'"hash": "{hsh}"' in html, "[fonte] 3D (index.html) embute o mesmo hash")
for a in dj["ambientes"]:
    ok(abs(a["area_bruta"] - idx[a["id"]]["area_bruta"]) < 1e-9, f"[fonte] 3D {a['id']} = CSV")
for f in ("02_planta_terreo.dxf", "03_planta_superior.dxf"):
    ok(os.path.getmtime(os.path.join(P.BASE, f)) >= os.path.getmtime(os.path.join(P.DADOS, "ambientes.csv")),
       f"[fonte] {f} mais novo que o CSV")

# 10. tabelas xlsx recalculadas
soffice = shutil.which("soffice") or shutil.which("libreoffice")
if soffice:
    tmp = tempfile.mkdtemp()
    for f in ("04_tabela_areas.xlsx", "10_orcamento.xlsx"):
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp,
                        os.path.join(P.BASE, f)], capture_output=True, timeout=180)
    wb = load_workbook(os.path.join(tmp, "04_tabela_areas.xlsx"), data_only=True)
    ws = wb["Meta"]
    ok(ws["B2"].value == hsh, "[xlsx] tabela de areas gerada com o hash atual")
    rs = wb["Resumo"]
    for row, p in ((4, "T"), (5, "S")):
        ok(abs(rs.cell(row=row, column=2).value - res[p]["bruta"]) < 1e-6, f"[xlsx] formula soma pav {p} = Python")
        ok(rs.cell(row=row, column=8).value == "OK", f"[xlsx] status de fechamento pav {p} = OK")
    ok(rs.cell(row=6, column=8).value == "OK", "[xlsx] total construido fecha")
    lo = wb["Lote"]
    ok(lo["B8"].value == "OK", "[xlsx] fechamento do lote = OK")
    wo = load_workbook(os.path.join(tmp, "10_orcamento.xlsx"), data_only=True)["Orcamento"]
    chk = [wo.cell(row=r, column=4).value for r in range(1, 60) if wo.cell(row=r, column=2).value and
           str(wo.cell(row=r, column=2).value).startswith("Conferencia")]
    ok(chk == ["OK"], "[orcamento] percentuais das categorias somam 100%")
else:
    avisos.append("[xlsx] LibreOffice nao encontrado - formulas nao recalculadas")

# 11. avisos de medida aproximada nos documentos
for f in sorted(os.listdir(P.BASE)):
    if f.endswith(".md"):
        t = open(os.path.join(P.BASE, f), encoding="utf-8").read().lower()
        ok("aproximad" in t or "conceitua" in t, f"[docs] {f} declara medidas aproximadas/conceituais")

print(f"OK: {len(oks)} | AVISOS: {len(avisos)} | FALHAS: {len(falhas)}")
for a in avisos:
    print("  AVISO :", a)
for f in falhas:
    print("  FALHA :", f)
sys.exit(1 if falhas else 0)
