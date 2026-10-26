"""Gera 10_orcamento.xlsx - estimativa parametrica em tres cenarios (ESTIMATIVA, nao e orcamento de obra).

Metodo: custo base = area equivalente (04_tabela_areas, NBR 12721) x CUB-SP do padrao do cenario,
distribuido em categorias por percentuais tipicos (HIPOTESE). Itens que o CUB NAO inclui
(projetos/taxas, fundacoes, churrasqueira/exaustao, luminarias decorativas, paisagismo) entram como
verbas separadas, para nao contar duas vezes. Administracao/BDI e contingencia vem abaixo das 22 categorias.
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import projeto as P

AZUL = Font(color="0000FF")
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="D9D2C5")
INP = PatternFill("solid", fgColor="FFF8E1")
TOT = PatternFill("solid", fgColor="EDE7DC")
THIN = Border(*(Side(style="thin", color="BBBBBB"),) * 4)
BRL = 'R$ #,##0;[Red]-R$ #,##0'

CUB = {  # SindusCon-SP, ref. agosto/2026, R$/m2, projeto-padrao R-1 (reproducao Portal DORI 03/09/2026)
    "Economico": (2190.76, "R1-B (baixo)"),
    "Intermediario": (2686.65, "R1-N (normal)"),
    "Alto padrao": (3249.37, "R1-A (alto)"),
}
CEN = list(CUB)

# (n, categoria, base, %CUB ou verba por cenario, % material, observacao)
CATS = [
    (1, "Projetos (arq., estrutural, instalacoes), aprovacao, ART/RRT e taxas", "verba", (18000, 30000, 50000), 0.00,
     "Fora do CUB. Sondagem e topografia estao na cat. 2"),
    (2, "Servicos preliminares: sondagem, topografia, canteiro, ligacoes provisorias", "cub", 0.02, 0.35, ""),
    (3, "Movimento de terra, locacao e regularizacao do terreno", "cub", 0.01, 0.10, "Depende da topografia (pendencia)"),
    (4, "Fundacoes (sapatas/estacas + baldrame)", "verba", (30000, 38000, 45000), 0.55,
     "Fora do CUB. Valor DEPENDE DA SONDAGEM - pode variar muito"),
    (5, "Estrutura complementar (pilares, vigas, lajes; vigas de 7 m sobre a garagem)", "cub", 0.18, 0.60,
     "Inclui reforco da laje da sacada-jardim"),
    (6, "Paineis de EPS - fornecimento e montagem (vedacoes)", "cub", 0.09, 0.65, "Substitui alvenaria"),
    (7, "Revestimento dos paineis (argamassa/microconcreto projetado 2 faces)", "cub", 0.06, 0.45, ""),
    (8, "Cobertura: laje, platibanda, telhado embutido termoacustico, calhas, claraboia", "cub", 0.05, 0.60, ""),
    (9, "Impermeabilizacao (baldrame, areas molhadas, sacada-jardim, lajes, jardineiras)", "cub", 0.04, 0.55, ""),
    (10, "Instalacoes hidrossanitarias, aguas pluviais e gas", "cub", 0.08, 0.55, ""),
    (11, "Instalacoes eletricas, dados/telecom, quadros e circuitos", "cub", 0.07, 0.55,
     "Infraestrutura e acabamentos basicos; luminarias decorativas na cat. 21"),
    (12, "Esquadrias de aluminio preto fosco e vidros (guarda-corpos na cat. 19)", "cub", 0.09,
     0.80, ""),
    (13, "Portas internas, porta pivotante de madeira e ferragens", "cub", 0.03, 0.75, ""),
    (14, "Contrapiso e revestimentos de piso (porcelanatos, amadeirado, externo, garagem)", "cub", 0.08, 0.60, ""),
    (15, "Revestimentos de parede (areas molhadas, lavabo, parede de destaque)", "cub", 0.03, 0.60, ""),
    (16, "Forros de gesso, sancas e tabicas", "cub", 0.04, 0.45, ""),
    (17, "Pintura e texturas (interna e fachada)", "cub", 0.05, 0.35, ""),
    (18, "Loucas, metais e bancadas de pedra (granito/quartzo)", "cub", 0.05, 0.80, ""),
    (19, "Escada (estrutura leve + degraus amadeirados), guarda-corpos e corrimaos", "cub", 0.03, 0.70, ""),
    (20, "Area gourmet: churrasqueira embutida, exaustao, bancada e revestimentos", "verba", (8000, 15000, 28000),
     0.65, "Fora do CUB (equipamentos)"),
    (21, "Iluminacao: luminarias, spots, perfis LED, fachada, jardim e sacada", "verba", (6000, 14000, 25000), 0.80,
     "Fora do CUB (luminarias decorativas)"),
    (22, "Paisagismo, jardineiras (vegetacao e substrato) e limpeza final", "verba", (5000, 10000, 18000), 0.55,
     "Fora do CUB (ajardinamento)"),
]
ADM = (0.05, 0.12, 0.15)     # administracao/BDI: economico = administracao propria
CONT = (0.10, 0.10, 0.10)    # contingencia
FORA = ["Moveis planejados completos (cozinha, closet, dormitorios, home, banheiros)", "Eletrodomesticos",
        "Muro frontal e muros de divisa", "Portao(oes) do lote", "Calcada publica", "Decoracao, cortinas e tapetes",
        "Ar-condicionado e aquecimento solar/gas (equipamentos)", "Terreno, ITBI, escritura e financiamento",
        "Ligacoes definitivas cobradas pelas concessionarias", "Mudanca de padrao do cliente durante a obra"]


def gerar():
    amb = P.ambientes()
    res = P.resumo(amb)
    area_constr = sum(v["bruta"] for v in res.values())
    area_eq = sum(a["area_bruta"] * P.COEF_EQUIV.get(a["id"], 1.0 if a["computa"] else 0.0) for a in amb)

    wb = Workbook()
    ws = wb.active
    ws.title = "Orcamento"
    ws["A1"] = "ORCAMENTO PRELIMINAR PARAMETRICO - 3 CENARIOS (ESTIMATIVA, NAO E ORCAMENTO DE OBRA)"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Base: CUB-SP (SindusCon-SP) ref. ago/2026 x area equivalente; percentuais por categoria e verbas = "
                "HIPOTESES do estudo. Valida-se somente com 3 orcamentos locais (Cajamar / regiao) por etapa.")
    # premissas
    ws["A4"], ws["B4"] = "Area construida (m2) - 04_tabela_areas", round(area_constr, 2)
    ws["A5"], ws["B5"] = "Area equivalente de custo (m2) - NBR 12721, coef. hipoteticos", round(area_eq, 2)
    for c in ("B4", "B5"):
        ws[c].font, ws[c].fill = AZUL, INP
    ws["A7"] = "Parametro"
    for j, c in enumerate(CEN):
        ws.cell(row=7, column=2 + j, value=c)
    for c in range(1, 5):
        ws.cell(row=7, column=c).font, ws.cell(row=7, column=c).fill = BOLD, HEAD
    ws["A8"] = "CUB-SP ago/2026 (R$/m2) - padrao"
    ws["A9"] = "Padrao CUB adotado"
    ws["A10"] = "Custo base CUB = area equiv. x CUB (R$)"
    ws["A11"] = "Administracao / BDI (% sobre subtotal)"
    ws["A12"] = "Contingencia (% sobre subtotal + adm.)"
    for j, c in enumerate(CEN):
        col = get_column_letter(2 + j)
        ws[f"{col}8"], ws[f"{col}9"] = CUB[c][0], CUB[c][1]
        ws[f"{col}8"].font, ws[f"{col}8"].fill = AZUL, INP
        ws[f"{col}8"].number_format = '#,##0.00'
        ws[f"{col}10"] = f"=$B$5*{col}8"
        ws[f"{col}10"].number_format = BRL
        ws[f"{col}11"], ws[f"{col}12"] = ADM[j], CONT[j]
        for r in (11, 12):
            ws[f"{col}{r}"].font, ws[f"{col}{r}"].fill, ws[f"{col}{r}"].number_format = AZUL, INP, "0%"

    # tabela de categorias
    hr = 14
    heads = ["#", "Categoria", "Base", "% do CUB (se base=cub)", "% material (hipotese)"]
    for c in CEN:
        heads += [f"{c} - materiais", f"{c} - servicos", f"{c} - total"]
    heads += ["Observacao"]
    for i, h in enumerate(heads, 1):
        cell = ws.cell(row=hr, column=i, value=h)
        cell.font, cell.fill, cell.border = BOLD, HEAD, THIN
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    widths = [4, 62, 7, 9, 9] + [14, 14, 14] * 3 + [50]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[hr].height = 42
    r = hr
    for n, nome, base, val, pm, obs in CATS:
        r += 1
        ws.cell(row=r, column=1, value=n)
        ws.cell(row=r, column=2, value=nome)
        ws.cell(row=r, column=3, value=base)
        if base == "cub":
            c = ws.cell(row=r, column=4, value=val)
            c.font, c.fill, c.number_format = AZUL, INP, "0%"
        c = ws.cell(row=r, column=5, value=pm)
        c.font, c.fill, c.number_format = AZUL, INP, "0%"
        for j in range(3):
            cm, cs, ct = 6 + 3 * j, 7 + 3 * j, 8 + 3 * j
            colcub = get_column_letter(2 + j)
            if base == "cub":
                ws.cell(row=r, column=ct, value=f"=$D{r}*{colcub}$10")
            else:
                c = ws.cell(row=r, column=ct, value=val[j])
                c.font, c.fill = AZUL, INP
            ws.cell(row=r, column=cm, value=f"={get_column_letter(ct)}{r}*$E{r}")
            ws.cell(row=r, column=cs, value=f"={get_column_letter(ct)}{r}-{get_column_letter(cm)}{r}")
            for cc in (cm, cs, ct):
                ws.cell(row=r, column=cc).number_format = BRL
        ws.cell(row=r, column=6 + 9, value=obs)
        for cc in range(1, 16):
            ws.cell(row=r, column=cc).border = THIN
    r1, r2 = hr + 1, r
    lin = {}
    for rot, chave in (("Subtotal das 22 categorias", "sub"), ("Administracao / BDI", "adm"),
                       ("Contingencia", "cont"), ("TOTAL ESTIMADO", "tot"), ("Custo por m2 construido (R$/m2)", "m2"),
                       ("Faixa do briefing - minimo (R$)", "bmin"), ("Faixa do briefing - maximo (R$)", "bmax"),
                       ("Diferenca total x maximo do briefing (R$)", "dif"), ("Situacao", "sit"),
                       ("Conferencia: soma % CUB das categorias = 100%?", "chk")):
        r += 1
        lin[chave] = r
        ws.cell(row=r, column=2, value=rot).font = BOLD
        for cc in range(1, 16):
            ws.cell(row=r, column=cc).fill = TOT
    for j in range(3):
        colcub = get_column_letter(2 + j)
        for k in range(3):
            col = get_column_letter(6 + 3 * j + k)
            ws[f"{col}{lin['sub']}"] = f"=SUM({col}{r1}:{col}{r2})"
            ws[f"{col}{lin['sub']}"].number_format = BRL
        ct = get_column_letter(8 + 3 * j)
        ws[f"{ct}{lin['adm']}"] = f"={ct}{lin['sub']}*{colcub}11"
        ws[f"{ct}{lin['cont']}"] = f"=({ct}{lin['sub']}+{ct}{lin['adm']})*{colcub}12"
        ws[f"{ct}{lin['tot']}"] = f"={ct}{lin['sub']}+{ct}{lin['adm']}+{ct}{lin['cont']}"
        ws[f"{ct}{lin['m2']}"] = f"={ct}{lin['tot']}/$B$4"
        ws[f"{ct}{lin['bmin']}"] = 360000
        ws[f"{ct}{lin['bmax']}"] = 480000
        ws[f"{ct}{lin['dif']}"] = f"={ct}{lin['tot']}-{ct}{lin['bmax']}"
        ws[f"{ct}{lin['sit']}"] = (f'=IF({ct}{lin["tot"]}<={ct}{lin["bmax"]},"DENTRO DA FAIXA",'
                                    f'"ACIMA DA FAIXA")')
        for k in ("adm", "cont", "tot", "m2", "bmin", "bmax", "dif"):
            ws[f"{ct}{lin[k]}"].number_format = BRL
        ws[f"{ct}{lin['tot']}"].font = BOLD
    ws[f"D{lin['chk']}"] = f'=IF(ABS(SUM(D{r1}:D{r2})-1)<0.0001,"OK","FALHA")'
    ws.freeze_panes = f"C{hr + 1}"

    # ---------------------------------------------------------------- fora do escopo
    fo = wb.create_sheet("Fora do escopo")
    fo["A1"] = "ITENS NAO INCLUIDOS NO TOTAL (orcar a parte)"
    fo["A1"].font = BOLD
    for i, t in enumerate(FORA, 3):
        fo.cell(row=i, column=1, value=t)
    fo.column_dimensions["A"].width = 90

    # ---------------------------------------------------------------- fontes
    fs = wb.create_sheet("Fontes")
    fontes = [
        ("Referencia", "Valor", "Data de referencia", "Fonte", "Status"),
        ("CUB-SP R1-B (residencia unifamiliar, padrao baixo)", "R$ 2.190,76/m2", "ago/2026",
         "SindusCon-SP, reproduzido por Portal DORI em 03/09/2026: https://portaldori.com.br/2026/09/03/"
         "sinduscon-divulga-tabelas-de-custo-unitario-basico-no-estado-de-sao-paulo-referencia-agosto-de-2026/",
         "Fonte secundaria; conferir no site oficial do SindusCon-SP"),
        ("CUB-SP R1-N (padrao normal)", "R$ 2.686,65/m2", "ago/2026", "idem", "Fonte secundaria"),
        ("CUB-SP R1-A (padrao alto)", "R$ 3.249,37/m2", "ago/2026", "idem", "Fonte secundaria"),
        ("SINAPI - custo medio m2 Estado de SP", "R$ 2.037,20/m2", "mai/2026",
         "IBGE, caderno SINAPI 05/2026: https://ftp.ibge.gov.br/Precos_Custos_e_Indices_da_Construcao_Civil/"
         "Fasciculo_Indicadores_IBGE/sinapi_202605caderno.pdf", "Referencia cruzada (padrao popular, sem acabamentos superiores)"),
        ("SINAPI - custo medio m2 Brasil", "R$ 1.999,73/m2", "set/2026",
         "IBGE, caderno SINAPI 09/2026: https://ftp.ibge.gov.br/Precos_Custos_e_Indices_da_Construcao_Civil/"
         "Fasciculo_Indicadores_IBGE/sinapi_202609caderno.pdf", "Referencia cruzada"),
        ("Distribuicao % por categoria, % material, verbas, ADM, contingencia", "-", "out/2026",
         "Estimativa do estudo (sem fonte verificavel)", "ESTIMATIVA - validar com orcamentos locais"),
        ("Coeficientes de area equivalente (garagem 0,75; gourmet 0,75; sacada descoberta 0,40)", "-", "-",
         "Faixas usuais da ABNT NBR 12721 - valores adotados sao hipotese", "HIPOTESE"),
        ("Observacao", "O CUB nao inclui fundacoes, projetos, taxas, BDI, paisagismo, equipamentos e itens especiais "
         "(NBR 12721); por isso esses itens entram como verbas separadas.", "", "", ""),
        ("Observacao", "Os valores do CUB sao do Estado de SP (amostra da capital/regiao). Cajamar (RMSP) tende a ter "
         "precos proximos, mas mao de obra especializada em EPS pode ser escassa e mais cara.", "", "", ""),
    ]
    for i, row in enumerate(fontes, 1):
        for j, v in enumerate(row, 1):
            c = fs.cell(row=i, column=j, value=v)
            if i == 1:
                c.font, c.fill = BOLD, HEAD
    for col, w in zip("ABCDE", (60, 18, 16, 100, 45)):
        fs.column_dimensions[col].width = w

    out = os.path.join(P.BASE, "10_orcamento.xlsx")
    wb.save(out)
    return out


if __name__ == "__main__":
    print(gerar())
