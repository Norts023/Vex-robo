"""Gera 04_tabela_areas.xlsx com formulas (soma dos ambientes x area do pavimento)."""
import datetime
import os

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import projeto as P

AZUL = Font(color="0000FF")          # entradas (dados do CSV)
PRETO = Font(color="000000")         # formulas
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="D9D2C5")
INPUT_BG = PatternFill("solid", fgColor="FFF8E1")
THIN = Border(*(Side(style="thin", color="BBBBBB"),) * 4)
VERDE = PatternFill("solid", fgColor="C8E6C9")
VERM = PatternFill("solid", fgColor="FFCDD2")


def cab(ws, row, titles, widths=None):
    for i, t in enumerate(titles, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill, c.border = BOLD, HEAD, THIN
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def gerar():
    amb = P.ambientes()
    pav = P.pavimentos()
    hsh = P.hash_dados()
    wb = Workbook()

    # ------------------------------------------------------------ Ambientes
    ws = wb.active
    ws.title = "Ambientes"
    ws["A1"] = "TABELA DE AREAS - ESTUDO CONCEITUAL (medidas aproximadas, nao executivas)"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = (f"Fonte unica: dados/ambientes.csv (hash {hsh}). Azul = dado de entrada; preto = formula. "
                f"Parede conceitual {P.T_PAREDE:.2f} m: lado externo desconta 0,15; lado interno 0,075; lado aberto 0.")
    titles = ["ID", "Pav.", "Ambiente", "Tipo", "Computa (1/0)", "x0", "y0", "x1", "y1",
              "Largura eixo (m)", "Prof. eixo (m)", "Area bruta (m2)", "Desc. W", "Desc. E", "Desc. S", "Desc. N",
              "Largura util (m)", "Prof. util (m)", "Area util (m2)", "Ref. briefing (m2)", "Ref. (texto)",
              "Desvio util x ref.", "Coef. equiv. custo", "Area equiv. (m2)", "Observacao"]
    cab(ws, 4, titles, [9, 5, 34, 11, 9, 6, 6, 6, 6, 9, 9, 10, 7, 7, 7, 7, 9, 9, 10, 10, 22, 10, 9, 10, 60])
    r0 = 5
    for i, a in enumerate(amb):
        r = r0 + i
        coef = P.COEF_EQUIV.get(a["id"], 1.0 if a["computa"] else 0.0)
        vals = [a["id"], a["pav"], a["nome"], a["tipo"], a["computa"], a["x0"], a["y0"], a["x1"], a["y1"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border = AZUL, THIN
            if c >= 5:
                cell.fill = INPUT_BG
        ws.cell(row=r, column=10, value=f"=H{r}-F{r}")
        ws.cell(row=r, column=11, value=f"=I{r}-G{r}")
        ws.cell(row=r, column=12, value=f"=J{r}*K{r}")
        for c, lado in zip((13, 14, 15, 16), "WESN"):
            cell = ws.cell(row=r, column=c, value=a["ded"][lado])
            cell.font, cell.fill = AZUL, INPUT_BG
        ws.cell(row=r, column=17, value=f"=J{r}-M{r}-N{r}")
        ws.cell(row=r, column=18, value=f"=K{r}-O{r}-P{r}")
        ws.cell(row=r, column=19, value=f"=Q{r}*R{r}")
        c = ws.cell(row=r, column=20, value=a["area_ref"] or None)
        c.font = AZUL
        ws.cell(row=r, column=21, value=a["ref_txt"])
        ws.cell(row=r, column=22, value=f'=IF(AND(T{r}>0,E{r}=1),S{r}/T{r}-1,"")')
        c = ws.cell(row=r, column=23, value=coef)
        c.font, c.fill = AZUL, INPUT_BG
        ws.cell(row=r, column=24, value=f"=L{r}*W{r}")
        ws.cell(row=r, column=25, value=a["nota"])
        for c in range(10, 25):
            ws.cell(row=r, column=c).border = THIN
        for c in (10, 11, 12, 17, 18, 19, 24):
            ws.cell(row=r, column=c).number_format = "0.00"
        ws.cell(row=r, column=22).number_format = "+0%;-0%;0%"
    rl = r0 + len(amb) - 1
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:Y{rl}"

    # ------------------------------------------------------------ Resumo
    rs = wb.create_sheet("Resumo")
    rs["A1"] = "FECHAMENTO DE AREAS POR PAVIMENTO (criterio de aceite: desvio <= 2%)"
    rs["A1"].font = Font(bold=True, size=12)
    cab(rs, 3, ["Pavimento", "Soma bruta dos ambientes computaveis (m2)", "Soma util (m2)",
                "Paredes e estrutura (m2)", "Util + paredes (m2)", "Area declarada do pavimento (m2)",
                "Desvio |soma/declarada - 1|", "Status (tol. 2%)", "Fecha util+paredes=bruta?"],
        [22, 18, 14, 14, 14, 18, 14, 14, 16])
    rng = f"Ambientes!$B${r0}:$B${rl}"
    rngE = f"Ambientes!$E${r0}:$E${rl}"
    for i, (p, d) in enumerate(pav.items()):
        r = 4 + i
        rs.cell(row=r, column=1, value=d["nome"])
        rs.cell(row=r, column=2, value=f'=SUMIFS(Ambientes!$L${r0}:$L${rl},{rng},"{p}",{rngE},1)')
        rs.cell(row=r, column=3, value=f'=SUMIFS(Ambientes!$S${r0}:$S${rl},{rng},"{p}",{rngE},1)')
        rs.cell(row=r, column=4, value=f"=B{r}-C{r}")
        rs.cell(row=r, column=5, value=f"=C{r}+D{r}")
        c = rs.cell(row=r, column=6, value=d["area_declarada_m2"])
        c.font, c.fill = AZUL, INPUT_BG
        rs.cell(row=r, column=7, value=f"=ABS(B{r}/F{r}-1)")
        rs.cell(row=r, column=8, value=f'=IF(G{r}<=0.02,"OK","FALHA")')
        rs.cell(row=r, column=9, value=f'=IF(ABS(E{r}-B{r})<0.001,"OK","FALHA")')
        for c in range(2, 7):
            rs.cell(row=r, column=c).number_format = "0.00"
        rs.cell(row=r, column=7).number_format = "0.00%"
    rt = 4 + len(pav)
    rs.cell(row=rt, column=1, value="AREA CONSTRUIDA TOTAL").font = BOLD
    for c, L in ((2, "B"), (3, "C"), (4, "D"), (5, "E"), (6, "F")):
        rs.cell(row=rt, column=c, value=f"=SUM({L}4:{L}{rt - 1})").font = BOLD
        rs.cell(row=rt, column=c).number_format = "0.00"
    rs.cell(row=rt, column=7, value=f"=ABS(B{rt}/F{rt}-1)").number_format = "0.00%"
    rs.cell(row=rt, column=8, value=f'=IF(G{rt}<=0.02,"OK","FALHA")')
    for col in "HI":
        rs.conditional_formatting.add(f"{col}4:{col}{rt}", CellIsRule(operator="equal", formula=['"OK"'], fill=VERDE))
        rs.conditional_formatting.add(f"{col}4:{col}{rt}", CellIsRule(operator="equal", formula=['"FALHA"'], fill=VERM))

    r = rt + 2
    rs.cell(row=r, column=1, value="Indicadores").font = BOLD
    linhas = [
        ("Area equivalente de custo (m2) - NBR 12721 (coef. hipoteticos)", f"=SUM(Ambientes!$X${r0}:$X${rl})"),
        ("Area nao computada: sacada descoberta (m2)", f'=SUMIFS(Ambientes!$L${r0}:$L${rl},Ambientes!$A${r0}:$A${rl},"SACADA")'),
        ("Area nao computada: poco de luz / patio (m2)", f'=SUMIFS(Ambientes!$L${r0}:$L${rl},Ambientes!$A${r0}:$A${rl},"POCO")'),
        ("Meta do briefing - area construida minima (m2)", 125),
        ("Meta do briefing - area construida maxima (m2)", 140),
        ("Excedente sobre a meta maxima (m2)", f"=B{rt}-B{r + 5}"),
        ("Area construida sem garagem e gourmet (m2) - para comparar com programa 'de morar'",
         f'=B{rt}-SUMIFS(Ambientes!$L${r0}:$L${rl},Ambientes!$A${r0}:$A${rl},"GAR")'
         f'-SUMIFS(Ambientes!$L${r0}:$L${rl},Ambientes!$A${r0}:$A${rl},"GOURMET")'),
    ]
    for j, (t, f) in enumerate(linhas, 1):
        rs.cell(row=r + j, column=1, value=t)
        c = rs.cell(row=r + j, column=2, value=f)
        c.number_format = "0.00"
        if isinstance(f, (int, float)):
            c.font, c.fill = AZUL, INPUT_BG
    rs.column_dimensions["A"].width = 62

    # ------------------------------------------------------------ Lote
    lt = wb.create_sheet("Lote")
    lt["A1"] = "OCUPACAO E PERMEABILIDADE DO LOTE (parametros municipais = HIPOTESES, nao confirmados)"
    lt["A1"].font = Font(bold=True, size=12)
    rows = [
        ("Largura do lote (m)", P.LOTE_L, True),
        ("Profundidade do lote (m)", P.LOTE_P, True),
        ("Area do lote (m2)", "=B3*B4", False),
        ("Projecao da edificacao no terreo (m2)",
         f'=SUMIFS(Ambientes!$L${r0}:$L${rl},{rng},"T",Ambientes!$D${r0}:$D${rl},"fechado")'
         f'+SUMIFS(Ambientes!$L${r0}:$L${rl},{rng},"T",Ambientes!$D${r0}:$D${rl},"coberto")', False),
        ("Areas descobertas do terreo (recuo, patio, jardim) (m2)",
         f'=SUMIFS(Ambientes!$L${r0}:$L${rl},{rng},"T",Ambientes!$D${r0}:$D${rl},"descoberto")', False),
        ("Fechamento: projecao + descobertas = lote?", '=IF(ABS(B6+B7-B5)<0.001,"OK","FALHA")', False),
        ("Taxa de ocupacao (TO) resultante", "=B6/B5", False),
        ("TO maxima - HIPOTESE (confirmar na LUOS de Cajamar)", P.PREMISSAS_LEGAIS["taxa_ocupacao_max"], True),
        ("Situacao TO", '=IF(B9<=B10,"ATENDE HIPOTESE","CONFLITO - ver 13_decisoes")', False),
        ("Area permeavel estimada (m2)", None, False),
        ("Taxa de permeabilidade (TP) resultante", "=B12/B5", False),
        ("TP minima - HIPOTESE (confirmar na LUOS de Cajamar)", P.PREMISSAS_LEGAIS["taxa_permeab_min"], True),
        ("Situacao TP", '=IF(B13>=B14,"ATENDE HIPOTESE","CONFLITO")', False),
        ("Recuo frontal adotado - HIPOTESE (m)", P.PREMISSAS_LEGAIS["recuo_frontal_m"], True),
        ("Afastamento minimo de janela a divisa (m) - Codigo Civil art. 1.301", P.PREMISSAS_LEGAIS["afast_janela_divisa_m"], True),
    ]
    for j, (t, v, inp) in enumerate(rows):
        rr = 3 + j
        lt.cell(row=rr, column=1, value=t)
        c = lt.cell(row=rr, column=2, value=v)
        if inp:
            c.font, c.fill = AZUL, INPUT_BG
    for rr in (9, 10, 13, 14):
        lt.cell(row=rr, column=2).number_format = "0.0%"
    for rr in (5, 6, 7, 12):
        lt.cell(row=rr, column=2).number_format = "0.00"
    lt.column_dimensions["A"].width = 66
    lt.column_dimensions["B"].width = 26
    er = 21
    lt.cell(row=er - 1, column=1, value="Superficies externas (dados/externas.csv)").font = BOLD
    cab_ext = ["ID", "Superficie", "x0", "y0", "x1", "y1", "Area (m2)", "Coef. permeavel (hipotese)",
               "Area permeavel (m2)", "Acabamento"]
    for i, t in enumerate(cab_ext, 1):
        c = lt.cell(row=er, column=i, value=t)
        c.font, c.fill = BOLD, HEAD
    ext = P.externas()
    for i, e in enumerate(ext):
        rr = er + 1 + i
        for c, v in enumerate([e["id"], e["nome"], e["x0"], e["y0"], e["x1"], e["y1"]], 1):
            lt.cell(row=rr, column=c, value=v).font = AZUL
        lt.cell(row=rr, column=7, value=f"=(E{rr}-C{rr})*(F{rr}-D{rr})").number_format = "0.00"
        c = lt.cell(row=rr, column=8, value=e["coef_perm"])
        c.font, c.fill = AZUL, INPUT_BG
        lt.cell(row=rr, column=9, value=f"=G{rr}*H{rr}").number_format = "0.00"
        lt.cell(row=rr, column=10, value=e["acabamento"])
    ef = er + len(ext)
    lt.cell(row=ef + 1, column=1, value="Total").font = BOLD
    lt.cell(row=ef + 1, column=7, value=f"=SUM(G{er + 1}:G{ef})").number_format = "0.00"
    lt.cell(row=ef + 1, column=9, value=f"=SUM(I{er + 1}:I{ef})").number_format = "0.00"
    lt["B12"] = f"=I{ef + 1}"
    lt.cell(row=ef + 2, column=1,
            value="Conferencia: superficies externas = descobertas do terreo?")
    lt.cell(row=ef + 2, column=2, value=f'=IF(ABS(G{ef + 1}-B7)<0.001,"OK","FALHA")')
    for c in ("B8", "B11", "B15", f"B{ef + 2}"):
        lt.conditional_formatting.add(c, CellIsRule(operator="equal", formula=['"OK"'], fill=VERDE))
        lt.conditional_formatting.add(c, CellIsRule(operator="equal", formula=['"ATENDE HIPOTESE"'], fill=VERDE))
        lt.conditional_formatting.add(c, CellIsRule(operator="notEqual", formula=['"OK"'], fill=VERM))

    # ------------------------------------------------------------ Meta
    mt = wb.create_sheet("Meta")
    for i, (k, v) in enumerate([
        ("Projeto", "Casa contemporanea em EPS - Cajamar/SP - ESTUDO CONCEITUAL"),
        ("Hash dos dados (CSV)", hsh),
        ("Gerado em", datetime.date.today().isoformat()),
        ("Gerador", "scripts/gerar_tabela.py"),
        ("Aviso", "Medidas aproximadas. Nao substitui projeto executivo, calculo estrutural, ART/RRT nem aprovacao."),
    ], 1):
        mt.cell(row=i, column=1, value=k).font = BOLD
        mt.cell(row=i, column=2, value=v)
    mt.column_dimensions["A"].width = 24
    mt.column_dimensions["B"].width = 100

    out = os.path.join(P.BASE, "04_tabela_areas.xlsx")
    wb.save(out)
    return out


if __name__ == "__main__":
    print(gerar())
